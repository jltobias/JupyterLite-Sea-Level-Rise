"""Optional LOCAL audit for authorized users. Never emits site IDs or coordinates."""
import argparse
from pathlib import Path
import csv
from collections import Counter

def main():
    ap=argparse.ArgumentParser(description=__doc__);ap.add_argument('workbook',type=Path);args=ap.parse_args()
    from openpyxl import load_workbook  # Optional: pip install openpyxl
    expected={}
    with (Path(__file__).resolve().parents[1]/'content/data/poster_aggregates.csv').open(encoding='utf-8') as f:
        for r in csv.DictReader(f):
            if r['scope']=='portfolio': expected[(r['scenario'],int(r['year']))]=int(r['potentially_flooded'])
    wb=load_workbook(args.workbook,read_only=True,data_only=True)
    sheet=wb['Climate-Sites-w-coords'];rows=sheet.iter_rows(values_only=True);headers=next(rows)
    mapping={'A':'RCP 2.6','B':'RCP 4.5','C':'RCP 8.5'}
    fields={f'{letter}{year%100:02d}':(scenario,year) for letter,scenario in mapping.items() for year in range(2030,2101,10)}
    missing=set(fields)|{'PEPFARSUP'};missing-=set(headers)
    if missing: raise ValueError(f'Required columns missing: {sorted(missing)}')
    counts=Counter();supported=0;unknown=0
    for row in rows:
        d=dict(zip(headers,row))
        if d['PEPFARSUP']!=1: continue
        supported+=1
        for field,key in fields.items():
            v=d[field]
            if v not in {'RED','ORANGE','YELLOW','NOFILL'}: unknown+=1
            counts[key]+=v in {'RED','ORANGE','YELLOW'}
    wb.close()
    mismatches=[key for key,value in expected.items() if counts[key]!=value]
    # Report only pass/fail, not unpublished aggregate values.
    print('Published denominator matches:',supported==25873)
    print('All 24 published portfolio totals match:',not mismatches)
    print('No unknown classifications:',unknown==0)
    if supported!=25873 or mismatches or unknown: raise SystemExit('Audit did not match the published data contract.')
if __name__=='__main__':main()
